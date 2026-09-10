import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def normalize(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value)
    without_marks = "".join(char for char in decomposed if not unicodedata.combining(char))
    return re.sub(r"[^\w]+", " ", without_marks.lower().replace("ё", "е")).strip()


def candidates(value: str) -> list[str]:
    normalized = normalize(value)
    results = [normalized]
    results.append(re.sub(r"^(г|город|city)\s+", "", normalized).strip())
    first_part = normalize(value.split(",")[0])
    results.append(re.sub(r"^(г|город|city)\s+", "", first_part).strip())

    if "санкт петербург" in normalized or normalized in {
        "спб",
        "spb",
        "питер",
        "saint petersburg",
        "st petersburg",
    }:
        results.append("санкт петербург")

    if normalized == "н новгород":
        results.append("нижнии новгород")

    for result in list(results):
        results.append(
            re.sub(
                r"\s+(ленинградская|ленинградскои|московская|московскои)\s+област[ьи].*$",
                "",
                result,
            ).strip()
        )

    return [result for result in dict.fromkeys(results) if result]


source_path = Path(sys.argv[1])
root = Path(__file__).resolve().parent.parent
city_index = json.loads((root / "server-data" / "city-index.json").read_text())
worksheet = load_workbook(source_path, read_only=True, data_only=True)["участники"]
rows = [
    str(row[7]).strip()
    for row in worksheet.iter_rows(min_row=2, values_only=True)
    if row[7] and str(row[7]).strip()
]

groups: Counter[int] = Counter()
unmatched: Counter[str] = Counter()

for raw_city in rows:
    place_index = next(
        (
            city_index["aliases"][candidate]
            for candidate in candidates(raw_city)
            if candidate in city_index["aliases"]
        ),
        None,
    )

    if place_index is None:
        unmatched[raw_city] += 1
    else:
        groups[place_index] += 1

country_codes = {city_index["places"][place_index][4] for place_index in groups}

print(
    json.dumps(
        {
            "participants": len(rows),
            "mapped_participants": sum(groups.values()),
            "cities": len(groups),
            "countries": len(country_codes),
            "unmatched": dict(unmatched),
        },
        ensure_ascii=False,
    )
)
